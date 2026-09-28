import pandas as pd
import os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Importamos la función que registra los logs.
# El archivo actualizar_logs.py debe estar en la misma
# carpeta que este archivo.
from actualizar_logs import registrar_log


# ============================================================
# RUTAS DE LOS ARCHIVOS
# ============================================================

# Ruta principal donde se encuentra la carpeta Archivos_Ventas.
# Si se mueve toda esta carpeta a otra ubicación, solo es
# necesario cambiar esta ruta.
CARPETA_PRINCIPAL = r"C:\Users\User\Documents\Archivos_Ventas"

# Carpeta donde se encuentra el archivo original.
CARPETA_DATOS = os.path.join(
    CARPETA_PRINCIPAL,
    "Datos Ventas"
)

# Carpeta donde se guardarán los archivos generados.
CARPETA_RESUMEN = os.path.join(
    CARPETA_PRINCIPAL,
    "Resumen de Ventas"
)

# Nombre del archivo original.
ARCHIVO_ENTRADA = os.path.join(
    CARPETA_DATOS,
    "datos_ventas.xlsx"
)

# Nombre del archivo actualizado.
ARCHIVO_ACTUALIZADO = os.path.join(
    CARPETA_RESUMEN,
    "datos_ventas_actualizado.xlsx"
)

# Nombre del archivo con los resúmenes.
ARCHIVO_RESUMEN = os.path.join(
    CARPETA_RESUMEN,
    "resumen_ventas.xlsx"
)

# Archivo donde se registran las ejecuciones del programa.
# Se guarda en la carpeta Logs.
ARCHIVO_LOGS = os.path.join(
    CARPETA_PRINCIPAL,
    "Logs",
    "logs_ejecucion.xlsx"
)


def cargar_datos(ruta):
    # Cargamos todas las columnas del archivo original.
    # De esta forma también conservamos ID_Cruce.
    df = pd.read_excel(ruta)

    return df


def procesar_datos(df):
    # Convertimos la columna Fecha a formato datetime
    # para poder trabajar con el año y el mes.
    df["Fecha"] = pd.to_datetime(
        df["Fecha"],
        errors="coerce"
    )

    # Completamos los valores faltantes de Total_Venta
    # utilizando Cantidad por Precio_Unitario.
    df["Total_Venta"] = df["Total_Venta"].fillna(
        df["Cantidad"] * df["Precio_Unitario"]
    )

    # Agregamos una nueva columna con el número del mes.
    df["Mes"] = df["Fecha"].dt.month

    # Movemos la columna Mes para que quede justo
    # después de la columna Fecha.
    columnas = list(df.columns)
    columnas.remove("Mes")

    posicion_fecha = columnas.index("Fecha")
    columnas.insert(posicion_fecha + 1, "Mes")

    df = df[columnas]

    return df


def ajustar_columnas(hoja):
    # Ajusta el ancho de cada columna según el texto
    # más largo que contenga (incluyendo el encabezado).
    for columna in hoja.columns:
        letra = get_column_letter(columna[0].column)

        largo_maximo = max(
            len(str(celda.value)) if celda.value is not None else 0
            for celda in columna
        )

        hoja.column_dimensions[letra].width = max(largo_maximo + 6, 14)


def guardar_archivo_actualizado(df, ruta_salida):
    # Guardamos todos los datos procesados en un nuevo archivo.
    # Se mantienen todas las columnas originales,
    # incluyendo ID_Cruce.
    with pd.ExcelWriter(
        ruta_salida,
        engine="openpyxl"
    ) as writer:
        df.to_excel(
            writer,
            sheet_name="Sheet1",
            index=False
        )

        # Ajustamos el tamaño de las columnas antes de cerrar el archivo.
        ajustar_columnas(writer.sheets["Sheet1"])


def dar_formato_hoja(hoja):
    # Da un formato uniforme a una hoja del resumen:
    # encabezado destacado, columnas ajustadas y datos alineados.

    # Estilos que se usarán en la hoja.
    fuente_encabezado = Font(bold=True, color="FFFFFF")
    relleno_encabezado = PatternFill("solid", fgColor="1F4E78")
    borde = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF")
    )

    # Formato del encabezado (primera fila).
    for celda in hoja[1]:
        celda.font = fuente_encabezado
        celda.fill = relleno_encabezado
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        celda.border = borde

    # Ajustamos el ancho de las columnas.
    ajustar_columnas(hoja)

    # Formato de los datos, columna por columna.
    # No se modifica el formato de los valores, solo
    # los bordes y la alineación.
    for columna in hoja.columns:
        nombre = columna[0].value

        for celda in columna[1:]:
            celda.border = borde

            if nombre == "Total_Ventas":
                # Los valores de ventas van alineados a la derecha.
                celda.alignment = Alignment(horizontal="right")
            elif nombre == "Mes":
                # El número del mes va centrado.
                celda.alignment = Alignment(horizontal="center")
            else:
                # El texto (vendedores) va a la izquierda.
                celda.alignment = Alignment(horizontal="left")

    # Dejamos el encabezado fijo al hacer scroll.
    hoja.freeze_panes = "A2"


def generar_resumen(df, ruta_salida):
    # Filtramos únicamente las ventas correspondientes al año 2023.
    ventas_2023 = df[
        df["Fecha"].dt.year == 2023
    ].copy()

    # Calculamos el total de ventas por vendedor.
    resumen_ventas = (
        ventas_2023
        .groupby(
            "Vendedor",
            as_index=False
        )["Total_Venta"]
        .sum()
        .rename(
            columns={
                "Total_Venta": "Total_Ventas"
            }
        )
    )

    # Calculamos el total de ventas por mes.
    ventas_mensuales = (
        ventas_2023
        .groupby(
            "Mes",
            as_index=False
        )["Total_Venta"]
        .sum()
        .rename(
            columns={
                "Total_Venta": "Total_Ventas"
            }
        )
    )

    # Creamos el archivo de resumen con las dos hojas solicitadas.
    with pd.ExcelWriter(
        ruta_salida,
        engine="openpyxl"
    ) as writer:

        resumen_ventas.to_excel(
            writer,
            sheet_name="Resumen_Ventas",
            index=False
        )

        ventas_mensuales.to_excel(
            writer,
            sheet_name="Ventas_Mensuales",
            index=False
        )

        # Aplicamos el formato a cada hoja antes de cerrar el archivo.
        dar_formato_hoja(writer.sheets["Resumen_Ventas"])
        dar_formato_hoja(writer.sheets["Ventas_Mensuales"])


def main():
    try:
        # Verificamos que exista el archivo de entrada.
        # Si no existe, lo registramos como error en el log.
        if not os.path.exists(ARCHIVO_ENTRADA):
            mensaje = f"No se encontró el archivo: {ARCHIVO_ENTRADA}"
            print(mensaje)
            registrar_log(ARCHIVO_LOGS, "ERROR", mensaje)
            return

        # Creamos la carpeta de salida si todavía no existe.
        os.makedirs(
            CARPETA_RESUMEN,
            exist_ok=True
        )

        # Cargamos el archivo original.
        df = cargar_datos(ARCHIVO_ENTRADA)

        # Procesamos los datos.
        df = procesar_datos(df)

        # Guardamos una versión actualizada del archivo original.
        guardar_archivo_actualizado(
            df,
            ARCHIVO_ACTUALIZADO
        )

        # Generamos el archivo con los resúmenes de 2023.
        generar_resumen(
            df,
            ARCHIVO_RESUMEN
        )

        print("Proceso realizado correctamente.")
        print(
            f"Archivo actualizado: {ARCHIVO_ACTUALIZADO}"
        )
        print(
            f"Archivo de resumen: {ARCHIVO_RESUMEN}"
        )

        # Registramos la ejecución exitosa en el log,
        # indicando cuántas filas se procesaron.
        registrar_log(
            ARCHIVO_LOGS,
            "EXITOSO",
            f"Se procesaron {len(df)} filas."
        )

    except Exception as e:
        # Si ocurre cualquier error, lo mostramos en pantalla
        # y también lo guardamos en el log.
        print(f"Ocurrió un error durante el proceso: {e}")
        registrar_log(ARCHIVO_LOGS, "ERROR", str(e))


if __name__ == "__main__":
    main()