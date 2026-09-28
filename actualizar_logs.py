import os
from datetime import datetime
from openpyxl import Workbook, load_workbook


# ============================================================
# CONFIGURACIÓN DEL LOG
# ============================================================

# Encabezados que tendrá el archivo de logs.
ENCABEZADOS = ["Fecha", "Hora", "Estado", "Detalle"]

# Ancho de cada columna (A, B, C, D) para que se lea mejor.
ANCHOS_COLUMNAS = [12, 10, 10, 80]


def registrar_log(ruta_log, estado, detalle=""):
    # Agrega un registro al archivo de logs.
    # Si el archivo no existe, lo crea con sus encabezados.
    # Si ya existe, agrega una nueva fila al final sin borrar
    # los registros anteriores.
    try:
        # Creamos la carpeta del log si todavía no existe.
        os.makedirs(
            os.path.dirname(ruta_log),
            exist_ok=True
        )

        if os.path.exists(ruta_log):
            # Si el archivo ya existe, lo abrimos para
            # agregar el nuevo registro.
            libro = load_workbook(ruta_log)
            hoja = libro.active
        else:
            # Si no existe, creamos un archivo nuevo
            # con los encabezados.
            libro = Workbook()
            hoja = libro.active
            hoja.title = "Logs"
            hoja.append(ENCABEZADOS)

            # Ajustamos el ancho de las columnas.
            for columna, ancho in zip("ABCD", ANCHOS_COLUMNAS):
                hoja.column_dimensions[columna].width = ancho

        # Obtenemos la fecha y hora actuales.
        ahora = datetime.now()

        # Agregamos el nuevo registro al final de la hoja.
        hoja.append([
            ahora.strftime("%Y-%m-%d"),
            ahora.strftime("%H:%M:%S"),
            estado,
            detalle
        ])

        # Guardamos el archivo con el nuevo registro.
        libro.save(ruta_log)

    except PermissionError:
        # Ocurre si el archivo de logs está abierto en Excel.
        print(
            "No se pudo escribir el log: "
            "cierra el archivo de logs e intenta de nuevo."
        )

    except Exception as e:
        # Un fallo al escribir el log nunca debe
        # detener el programa principal.
        print(f"No se pudo registrar el log: {e}")